"""Generate an editable Word (.docx) version of the paper from paper.tex.

Why this exists: the user wants to edit the paper in Word rather than
LaTeX. Rather than hand-retype the text into a second document -- which
would let the two versions drift apart the moment either one is edited --
this parses paper.tex directly (a file we author and fully control, with a
small, known set of LaTeX commands) and converts it into a native, styled
Word document: real heading styles, real tables rebuilt from the same
.tex table sources, and the same figure images already committed under
paper/figures/. Every number in the resulting .docx traces back to
paper.tex, not to a second, independently-typed copy of the results.

This is a converter for OUR OWN authored document, not a general LaTeX
parser -- it only needs to understand the specific handful of commands
paper.tex actually uses (\\section, \\subsection, \\cite, \\ref, \\emph,
inline math, \\input of a table file, and a fixed figure/table shape).

Usage:
    python ml/scripts/make_paper_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

PAPER_DIR = Path("paper")
TEX = PAPER_DIR / "paper.tex"
OUT = PAPER_DIR / "paper.docx"

MATH_SUBS = [
    (r"\\bar\{x\}", "x\u0304"),
    (r"t_\{0\.975,\\,n-1\}", "t(0.975, n\u22121)"),
    (r"\\times", "\u00d7"),
    (r"\\pm", "\u00b1"),
    (r"\\cdot\s*", "\u00b7"),  # trailing space after \cdot is a LaTeX token separator, not a printed space
    (r"\\sqrt\{n\}", "\u221an"),
    (r"\\log\(p\)", "log(p)"),
    (r"\\alpha", "\u03b1"),
]


def load_bib_order(text: str) -> dict:
    keys = re.findall(r"\\bibitem\{([^}]+)\}", text)
    return {k: i + 1 for i, k in enumerate(keys)}


def load_bib_entries(text: str) -> list:
    block = text.split("\\begin{thebibliography}")[1].split("\\end{thebibliography}")[0]
    entries = re.split(r"\\bibitem\{[^}]+\}", block)[1:]
    return [inline_to_plain(e.strip().replace("\n", " ")) for e in entries]


# Bare number/roman-numeral each \ref{...} resolves to INLINE. Every
# \ref{} in paper.tex is already preceded by its own literal "Table~",
# "Fig.~", or "Section~" in the source text, so substituting the bare
# value here (not "Table II") avoids doubling that word.
REF_LABELS = {
    "tab:datasets": "I",
    "tab:brain_mri": "II",
    "tab:crossdata": "III",
    "tab:chest_xray": "IV",
    "tab:mammography": "V",
    "fig:seed_lottery": "1",
    "fig:leakage": "2",
    "fig:crossdataset": "3",
    "fig:acc_auc": "4",
    "sec:results": "V",
}

# Full caption-header prefix, used only when WE construct a table/figure
# caption ourselves (there is no preceding literal word to reuse there).
CAPTION_LABELS = {
    "tab:datasets": "Table I",
    "tab:brain_mri": "Table II",
    "tab:crossdata": "Table III",
    "tab:chest_xray": "Table IV",
    "tab:mammography": "Table V",
    "fig:seed_lottery": "Fig. 1",
    "fig:leakage": "Fig. 2",
    "fig:crossdataset": "Fig. 3",
    "fig:acc_auc": "Fig. 4",
}

# Backwards-compat alias so inline_to_plain (used for captions/bib entries,
# which never have a preceding literal word of their own) still resolves
# a bare \ref{} to something readable if one ever appears there.
LABELS = CAPTION_LABELS


def inline_to_plain(s: str) -> str:
    """Strip inline LaTeX down to plain text (used for refs/captions where
    we don't need run-level formatting, e.g. bibliography entries)."""
    for pat, rep in MATH_SUBS:
        s = re.sub(pat, rep, s)
    s = re.sub(r"\$([^$]*)\$", r"\1", s)
    s = re.sub(r"\\emph\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\cite\{[^}]+\}", "", s)
    s = re.sub(r"\\ref\{([^}]+)\}", lambda m: LABELS.get(m.group(1), m.group(1)), s)
    s = s.replace("{,}", ",")
    # \{ and \} are LaTeX for a LITERAL brace character (e.g. the set
    # notation {42,0,1,2,3}), distinct from a plain { } grouping pair --
    # protect them before the generic grouping-brace strip below removes
    # every bare { or } indiscriminately.
    s = s.replace("\\{", "\x00OB\x00").replace("\\}", "\x00CB\x00")
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    # LaTeX's escaped-special-character sequences: a literal %, _, &, #, $
    # in the source, not a command -- must be unescaped before the generic
    # "\\command" stripper below, or the backslash and letter both survive.
    for esc, lit in (("\\%", "%"), ("\\_", "_"), ("\\&", "&"), ("\\#", "#"), ("\\$", "$")):
        s = s.replace(esc, lit)
    s = s.replace("~", " ")
    s = s.replace("``", "\u201c").replace("''", "\u201d")
    s = re.sub(r"\\[a-zA-Z]+", "", s)
    s = re.sub(r"[{}]", "", s)
    s = s.replace("\x00OB\x00", "{").replace("\x00CB\x00", "}")
    return re.sub(r"\s+", " ", s).strip()


def add_runs(paragraph, text: str, bib_order: dict) -> None:
    """Walk the raw LaTeX paragraph text left to right, emitting Word runs
    with the right formatting for \\emph{}, $math$, \\cite{}, and \\ref{}
    spans, and plain runs for everything else."""
    token_re = re.compile(
        r"\\emph\{(?P<emph>[^}]*)\}"
        r"|\$(?P<math>[^$]*)\$"
        r"|\\cite\{(?P<cite>[^}]+)\}"
        r"|\\ref\{(?P<ref>[^}]+)\}"
    )
    pos = 0
    for m in token_re.finditer(text):
        if m.start() > pos:
            _plain_run(paragraph, text[pos:m.start()])
        if m.group("emph") is not None:
            r = paragraph.add_run(_clean_plain(m.group("emph")))
            r.italic = True
        elif m.group("math") is not None:
            content = m.group("math")
            for pat, rep in MATH_SUBS:
                content = re.sub(pat, rep, content)
            content = re.sub(r"[{}\\]", "", content)
            r = paragraph.add_run(content)
            r.italic = True
        elif m.group("cite") is not None:
            keys = [k.strip() for k in m.group("cite").split(",")]
            nums = [str(bib_order[k]) for k in keys]
            paragraph.add_run(f"[{','.join(nums)}]")
        elif m.group("ref") is not None:
            paragraph.add_run(REF_LABELS.get(m.group("ref"), m.group("ref")))
        pos = m.end()
    if pos < len(text):
        _plain_run(paragraph, text[pos:])


def _clean_plain(s: str) -> str:
    s = s.replace("{,}", ",")
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    for esc, lit in (("\\%", "%"), ("\\_", "_"), ("\\&", "&"), ("\\#", "#"), ("\\$", "$")):
        s = s.replace(esc, lit)
    s = s.replace("~", " ")
    s = s.replace("``", "\u201c").replace("''", "\u201d")
    return s


def _plain_run(paragraph, s: str) -> None:
    s = _clean_plain(s)
    if s:
        paragraph.add_run(s)


def parse_tex_table(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")
    caption = re.search(r"\\caption\{(.+?)\}\s*\n\\label", text, re.S).group(1)
    # Read the real \label{} rather than guessing it from the filename --
    # tab_crossdataset.tex's own \label is the shorter "tab:crossdata",
    # so reconstructing the key from the filename silently mismatches and
    # produces a caption with no table number at all.
    label = re.search(r"\\label\{([^}]+)\}", text).group(1)
    body = text.split("\\begin{tabular}")[1].split("\\end{tabular}")[0]
    body = body.split("}", 1)[1]  # drop the {lcc} column spec
    rows = []
    for line in body.splitlines():
        line = line.strip()
        if not line or line.startswith(("\\toprule", "\\midrule", "\\bottomrule")):
            continue
        line = line.rstrip("\\").strip()
        cells = [c.strip() for c in line.split("&")]
        rows.append(cells)
    return {"caption": inline_to_plain(caption), "rows": rows, "label": label}


def add_table(doc: Document, label: str, tex_data: dict) -> None:
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"{label}. {tex_data['caption']}")
    r.italic = True
    r.font.size = Pt(9)

    rows = tex_data["rows"]
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell_text = inline_to_plain(row[j]) if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = cell_text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
    doc.add_paragraph()


def add_figure(doc: Document, label: str, png_path: Path, caption: str) -> None:
    doc.add_picture(str(png_path), width=Inches(3.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"{label}. {caption}")
    r.italic = True
    r.font.size = Pt(9)
    doc.add_paragraph()


def main() -> None:
    text = TEX.read_text(encoding="utf-8")
    bib_order = load_bib_order(text)

    title = re.search(r"\\title\{(.+?)\}\s*\n", text, re.S).group(1)
    abstract = re.search(r"\\begin\{abstract\}(.+?)\\end\{abstract\}", text, re.S).group(1).strip()
    keywords = re.search(r"\\begin\{IEEEkeywords\}(.+?)\\end\{IEEEkeywords\}", text, re.S).group(1).strip()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    t = doc.add_heading(level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(t, title, bib_order)

    byline = doc.add_paragraph()
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline.add_run("Aditya Ayushman Sahoo\n").bold = True
    byline.add_run("Independent Researcher \u2014 please replace with your institution before submission\n")
    byline.add_run("adityaasahoo@gmail.com")

    doc.add_paragraph()
    h = doc.add_paragraph()
    h.add_run("Abstract\u2014").bold = True
    add_runs(h, abstract, bib_order)
    for r in h.runs[1:]:
        r.italic = True

    h2 = doc.add_paragraph()
    h2.add_run("Index Terms\u2014").bold = True
    idx = h2.add_run(inline_to_plain(keywords))
    idx.italic = True
    doc.add_paragraph()

    body = text.split("\\maketitle", 1)[1].split("\\begin{thebibliography}")[0]
    body = body.split("\\begin{IEEEkeywords}")[1].split("\\end{IEEEkeywords}", 1)[1]

    section_num = 0
    subsection_letter = 0
    line_iter = iter(body.split("\n\n"))
    for block in line_iter:
        block = block.strip()
        if not block:
            continue

        m = re.match(r"\\section\{(.+?)\}", block)
        if m:
            section_num += 1
            subsection_letter = 0
            roman = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII"][section_num]
            doc.add_heading(f"{roman}. {inline_to_plain(m.group(1))}", level=1)
            rest = block[m.end():].strip()
            if rest.startswith("\\label"):
                rest = re.sub(r"^\\label\{[^}]+\}", "", rest).strip()
            if rest:
                p = doc.add_paragraph()
                add_runs(p, rest, bib_order)
            continue

        m = re.match(r"\\subsection\{(.+?)\}", block)
        if m:
            subsection_letter += 1
            letter = chr(ord("A") + subsection_letter - 1)
            doc.add_heading(f"{letter}. {inline_to_plain(m.group(1))}", level=2)
            rest = block[m.end():].strip()
            if rest:
                p = doc.add_paragraph()
                add_runs(p, rest, bib_order)
            continue

        m = re.match(r"\\input\{tables/(.+?)\.tex\}", block)
        if m:
            tex_data = parse_tex_table(PAPER_DIR / "tables" / f"{m.group(1)}.tex")
            add_table(doc, CAPTION_LABELS.get(tex_data["label"], "Table"), tex_data)
            continue

        m = re.match(r"\\begin\{table\}.*?\\caption\{(.+?)\}\s*\\label\{([^}]+)\}\s*"
                     r"\\begin\{tabular\}\{[^}]*\}(.+?)\\end\{tabular\}\s*\\end\{table\}",
                     block, re.S)
        if m:
            caption, label_key, tbody = m.groups()
            rows = []
            for line in tbody.splitlines():
                line = line.strip()
                if not line or line.startswith(("\\toprule", "\\midrule", "\\bottomrule")):
                    continue
                line = line.rstrip("\\").strip()
                if line:
                    rows.append([c.strip() for c in line.split("&")])
            add_table(doc, LABELS.get(label_key, "Table"), {"caption": inline_to_plain(caption), "rows": rows})
            continue

        m = re.match(
            r"\\begin\{figure\}.*?\\includegraphics\[[^\]]*\]\{figures/([^}]+)\.pdf\}"
            r"\s*\\caption\{(.+?)\}\s*\\label\{([^}]+)\}",
            block, re.S,
        )
        if m:
            fname, caption, label_key = m.groups()
            png = PAPER_DIR / "figures" / f"{fname}.png"
            add_figure(doc, LABELS.get(label_key, "Fig."), png, inline_to_plain(caption))
            continue

        p = doc.add_paragraph()
        add_runs(p, block, bib_order)

    doc.add_heading("References", level=1)
    for i, entry in enumerate(load_bib_entries(text), start=1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(entry).font.size = Pt(10)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT} ({OUT.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
